import io

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from models import BoundingBox, QuestionRegion
from ppt.exporter import export_subject_ppts


def make_template() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9), Inches(0.5))
    title.text_frame.text = "#QUESTION | #SUBJECT | Answer: #ANSWER"
    image_area = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(8), Inches(0.3))
    image_area.name = "QUESTION_IMAGE"
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def test_template_based_ppt_generation_creates_nonempty_ppt():
    image = Image.new("RGB", (800, 500), "white")
    region = QuestionRegion(
        number=1, page_index=0, column_index=0,
        box=BoundingBox(0, 0, 800, 500), image=image, subject="Mathematics",
    )
    outputs = export_subject_ppts(make_template(), [region], {1: "A"})
    assert set(outputs) == {"Mathematics"}
    assert len(outputs["Mathematics"]) > 1000
    prs = Presentation(io.BytesIO(outputs["Mathematics"]))
    assert len(prs.slides) == 1
    assert any(shape.shape_type == 13 for shape in prs.slides[0].shapes)
    text = " ".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Q1" in text and "Mathematics" in text and "A" in text


def test_answer_is_mapped_by_question_number_and_visible_on_same_slide():
    image = Image.new("RGBA", (800, 500), (0, 0, 0, 0))
    r1 = QuestionRegion(
        number=1, page_index=0, column_index=0,
        box=BoundingBox(0, 0, 800, 500), image=image, subject="Physics",
    )
    r2 = QuestionRegion(
        number=21, page_index=0, column_index=0,
        box=BoundingBox(0, 0, 800, 500), image=image, subject="Physics",
    )
    outputs = export_subject_ppts(make_template(), [r2, r1], {1: "A", 21: "17"})
    prs = Presentation(io.BytesIO(outputs["Physics"]))
    assert len(prs.slides) == 2
    texts = [" ".join(sh.text for sh in slide.shapes if sh.has_text_frame) for slide in prs.slides]
    # Slides are sorted by question number, and each answer belongs to its own Q.
    assert "Q1" in texts[0] and "A" in texts[0]
    assert "Q21" in texts[1] and "17" in texts[1]
